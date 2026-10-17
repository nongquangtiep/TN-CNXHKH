import docx
import re

doc = docx.Document('CHỦ NGHĨA XH-KH.docx')

with open('analysis_out.txt', 'w', encoding='utf-8') as out:
    out.write("=== PARAGRAPH ANALYSIS ===\n")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if re.search(r'^(BÀI|Bài|Chương|CHƯƠNG)\s+\d+', text, re.IGNORECASE):
            out.write(f"Lesson at P{i:04d}: {text}\n")
        elif re.search(r'^(Câu|CÂU)\s+\d+', text, re.IGNORECASE):
            # check if there's special instructions like (Chọn 2 đáp án), True/False, etc.
            out.write(f"  Q P{i:04d}: {text}\n")

    out.write("\n=== TABLES ANALYSIS ===\n")
    for t_idx, table in enumerate(doc.tables):
        out.write(f"\n--- Table {t_idx} ({len(table.rows)} rows, {len(table.columns)} cols) ---\n")
        for r_idx, row in enumerate(table.rows):
            cells_text = []
            for c_idx, cell in enumerate(row.cells):
                cell_runs_info = []
                for p in cell.paragraphs:
                    for r in p.runs:
                        mods = []
                        if r.bold: mods.append('B')
                        if r.underline: mods.append('U')
                        if r.font.color and r.font.color.rgb: mods.append(f'C:{r.font.color.rgb}')
                        if r.font.highlight_color: mods.append(f'H:{r.font.highlight_color}')
                        if mods:
                            cell_runs_info.append(f"[{','.join(mods)}: {r.text}]")
                cell_str = cell.text.strip().replace('\n', ' ')
                if cell_runs_info:
                    cell_str += f" ({' '.join(cell_runs_info)})"
                cells_text.append(cell_str)
            out.write(f"  R{r_idx}: {' | '.join(cells_text)}\n")

print("Analysis finished successfully!")
