import docx
import json

doc = docx.Document('CHỦ NGHĨA XH-KH.docx')
print('Paragraphs count:', len(doc.paragraphs))
print('Tables count:', len(doc.tables))

with open('dump_raw.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        runs_info = []
        for r in p.runs:
            mods = []
            if r.bold: mods.append('B')
            if r.underline: mods.append('U')
            if r.font.color and r.font.color.rgb: mods.append(f'C:{r.font.color.rgb}')
            if r.font.highlight_color: mods.append(f'H:{r.font.highlight_color}')
            if mods:
                runs_info.append(f"[{','.join(mods)}: {r.text}]")
        f.write(f"P{i:04d}: {text}\n")
        if runs_info:
            f.write(f"       FORMATS: {' '.join(runs_info)}\n")

    if doc.tables:
        f.write("\n--- TABLES ---\n")
        for t_idx, table in enumerate(doc.tables):
            f.write(f"Table {t_idx} ({len(table.rows)}x{len(table.columns)}):\n")
            for r_idx, row in enumerate(table.rows):
                row_texts = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                f.write(f"  R{r_idx}: {' | '.join(row_texts)}\n")

print("Saved raw dump to dump_raw.txt")
