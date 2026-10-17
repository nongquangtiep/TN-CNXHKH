import docx
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table
import re
import json
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('CHỦ NGHĨA XH-KH.docx')

elements = []
for child in doc.element.body:
    if isinstance(child, CT_P):
        p = Paragraph(child, doc)
        elements.append(('P', p))
    elif isinstance(child, CT_Tbl):
        t = Table(child, doc)
        elements.append(('T', t))

# Helper to check if a run has red color
def is_red(run):
    if run.font.color and run.font.color.rgb:
        hex_code = f"{run.font.color.rgb}".upper()
        if any(hex_code.startswith(prefix) for prefix in ['EE', 'FF', 'C0', 'ED', 'E5', 'CC', 'D9']):
            return True
    return False

# Group elements into questions
raw_questions = []
current_lesson = "Bài 1"
current_lesson_num = 1
current_q = None

for idx, (etype, elem) in enumerate(elements):
    if etype == 'P':
        text = elem.text.strip()
        if not text:
            continue
        lesson_match = re.search(r'^(BÀI|Bài|Chương|CHƯƠNG)\s+(\d+)', text, re.IGNORECASE)
        if lesson_match:
            current_lesson_num = int(lesson_match.group(2))
            current_lesson = f"Bài {current_lesson_num}"
            continue
        
        q_match = re.search(r'^(Câu|CÂU)\s+(\d+)\s*[:.]?\s*(.*)', text, re.IGNORECASE)
        if q_match:
            if current_q:
                raw_questions.append(current_q)
            current_q = {
                'lesson': current_lesson,
                'lesson_num': current_lesson_num,
                'q_num': int(q_match.group(2)),
                'title_full': text,
                'prompt': q_match.group(3).strip() if q_match.group(3).strip() else text,
                'elements': [],
                'start_idx': idx
            }
        else:
            if current_q:
                current_q['elements'].append(('P', elem))
    elif etype == 'T':
        if current_q:
            current_q['elements'].append(('T', elem))

if current_q:
    raw_questions.append(current_q)

# Process each question
parsed_questions = []
errors = []

for q_idx, q in enumerate(raw_questions):
    qid = q_idx + 1
    lesson = q['lesson']
    lesson_num = q['lesson_num']
    q_num = q['q_num']
    title = q['title_full']
    prompt = q['prompt']
    elems = q['elements']
    
    # Check tables
    tables = [elem for etype, elem in elems if etype == 'T']
    
    # Check title cues
    is_tf = bool(re.search(r'đúng\s*(hoặc|\/)\s*sai', title, re.IGNORECASE))
    is_drag_title = bool(re.search(r'kéo\s*thả|khớp\s*với|ghép|cột\s*tương\s*ứng', title, re.IGNORECASE))
    is_multi_title = bool(re.search(r'chọn\s*[2345]\s*(đáp\s*án|phương\s*án)', title, re.IGNORECASE))
    
    # 1. DRAG AND DROP WITH TABLE
    if tables:
        tbl = tables[0]
        first_row_cells = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
        
        is_categorize = False
        if any(re.search(r'1\)|2\)|Điều kiện chủ quan|Đặc trưng của dân tộc', c, re.IGNORECASE) for c in first_row_cells):
            is_categorize = True
        
        if is_categorize:
            cat1_title = first_row_cells[0]
            cat2_title = first_row_cells[1] if len(first_row_cells) > 1 else "Cột 2"
            cat1_items = []
            cat2_items = []
            all_items = []
            for r in tbl.rows[1:]:
                c0 = r.cells[0].text.strip().replace('\n', ' ')
                c1 = r.cells[1].text.strip().replace('\n', ' ') if len(r.cells) > 1 else ""
                if c0:
                    cat1_items.append(c0)
                    all_items.append({"text": c0, "category": cat1_title})
                if c1:
                    cat2_items.append(c1)
                    all_items.append({"text": c1, "category": cat2_title})
            
            parsed_questions.append({
                "id": qid,
                "lesson": lesson,
                "lesson_num": lesson_num,
                "q_num": q_num,
                "type": "drag_drop_categorize",
                "question": title,
                "instruction": "Kéo hoặc chọn từng mục vào đúng cột tương ứng",
                "categories": [
                    {"title": cat1_title, "items": cat1_items},
                    {"title": cat2_title, "items": cat2_items}
                ],
                "all_items": all_items
            })
            continue
        else:
            pairs = []
            options_pool = []
            for r in tbl.rows:
                c0 = r.cells[0].text.strip().replace('\n', ' ')
                c1 = r.cells[1].text.strip().replace('\n', ' ') if len(r.cells) > 1 else ""
                c0 = c0.replace('troChính sách an sinh xã hộing', 'trong')
                pairs.append({"left": c0, "right": c1})
                if c1:
                    options_pool.append(c1)
            
            parsed_questions.append({
                "id": qid,
                "lesson": lesson,
                "lesson_num": lesson_num,
                "q_num": q_num,
                "type": "drag_drop_match",
                "question": title,
                "instruction": "Kéo hoặc chọn đáp án đúng ghép vào vế tương ứng",
                "pairs": pairs,
                "options_pool": list(set(options_pool))
            })
            continue

    # 2. DRAG AND DROP WITHOUT TABLE
    if is_drag_title and not tables:
        pairs = []
        pool = []
        for etype, elem in elems:
            if etype == 'P':
                txt = elem.text.strip()
                if not txt:
                    continue
                bracket_match = re.search(r'\[(.*?)\]', txt)
                if bracket_match:
                    ans = bracket_match.group(1).strip()
                    left = txt.replace(f'[{ans}]', '______')
                    pairs.append({"left": left, "right": ans})
                    pool.append(ans)
                else:
                    red_texts = []
                    for r in elem.runs:
                        if is_red(r):
                            red_texts.append(r.text.strip())
                    red_str = ' '.join(red_texts).strip()
                    if red_str:
                        if red_str == txt:
                            pool.append(red_str)
                        else:
                            left = txt.replace(red_str, '______')
                            pairs.append({"left": left, "right": red_str})
                            pool.append(red_str)
        
        parsed_questions.append({
            "id": qid,
            "lesson": lesson,
            "lesson_num": lesson_num,
            "q_num": q_num,
            "type": "drag_drop_match",
            "question": title,
            "instruction": "Kéo hoặc chọn đáp án đúng ghép vào vế tương ứng",
            "pairs": pairs,
            "options_pool": list(set(pool))
        })
        continue

    # 3. TRUE / FALSE
    if is_tf:
        items = []
        p_list = [elem for etype, elem in elems if etype == 'P' and elem.text.strip()]
        
        i = 0
        while i < len(p_list):
            p = p_list[i]
            txt = p.text.strip()
            
            same_line_match = re.search(r'^(.*?)\s+(Đúng|Sai)\s*$', txt, re.IGNORECASE)
            if same_line_match:
                stmt = same_line_match.group(1).strip()
                ans_str = same_line_match.group(2).strip()
                items.append({
                    "statement": stmt,
                    "correct": ans_str.lower() == 'đúng'
                })
                i += 1
            else:
                stmt = txt
                if i + 1 < len(p_list):
                    next_p = p_list[i+1]
                    next_txt = next_p.text.strip()
                    if re.match(r'^(Đúng|Sai)', next_txt, re.IGNORECASE):
                        ans_str = next_txt
                        items.append({
                            "statement": stmt,
                            "correct": ans_str.lower().startswith('đúng')
                        })
                        i += 2
                    else:
                        errors.append(f"{lesson} Q{q_num} (TF): Expected Đúng/Sai after '{stmt}', got '{next_txt}'")
                        i += 1
                else:
                    errors.append(f"{lesson} Q{q_num} (TF): Statement without answer at end: '{stmt}'")
                    i += 1
        
        parsed_questions.append({
            "id": qid,
            "lesson": lesson,
            "lesson_num": lesson_num,
            "q_num": q_num,
            "type": "true_false",
            "question": title,
            "instruction": "Chọn Đúng hoặc Sai cho mỗi phát biểu sau",
            "items": items
        })
        continue

    # 4. SINGLE OR MULTIPLE CHOICE
    options = []
    correct_indices = []
    
    p_list = [elem for etype, elem in elems if etype == 'P' and elem.text.strip()]
    
    flattened_options = []
    for p in p_list:
        txt = p.text.strip()
        lines = [l.strip() for l in txt.split('\n') if l.strip()]
        if len(lines) > 1 and any(re.match(r'^[A-D]\s*[\.\:]', l) for l in lines):
            for l in lines:
                is_opt_red = False
                for r in p.runs:
                    if r.text and (r.text.strip() in l or l in r.text.strip()):
                        if is_red(r):
                            is_opt_red = True
                flattened_options.append((l, is_opt_red))
        else:
            is_opt_red = any(is_red(r) for r in p.runs)
            flattened_options.append((txt, is_opt_red))

    for idx_opt, (opt_txt, is_opt_red) in enumerate(flattened_options):
        label_match = re.match(r'^([A-D])\s*[\.\:]\s*(.*)', opt_txt)
        if label_match:
            formatted_text = f"{label_match.group(1)}. {label_match.group(2)}"
        else:
            letters = ['A', 'B', 'C', 'D', 'E', 'F']
            letter = letters[idx_opt] if idx_opt < len(letters) else f"{idx_opt+1}"
            formatted_text = f"{letter}. {opt_txt}"
        
        options.append(formatted_text)
        if is_opt_red:
            correct_indices.append(idx_opt)

    is_multi = len(correct_indices) > 1 or is_multi_title
    qtype = "multiple_choice" if is_multi else "single_choice"
    
    instruction = "Chọn 1 đáp án đúng nhất"
    if qtype == "multiple_choice":
        num_ans = len(correct_indices) if len(correct_indices) > 1 else 2
        instruction = f"Chọn {num_ans} đáp án đúng"

    if len(correct_indices) == 0:
        errors.append(f"{lesson} Q{q_num} ({qtype}): NO CORRECT ANSWER DETECTED! Options: {options}")

    parsed_questions.append({
        "id": qid,
        "lesson": lesson,
        "lesson_num": lesson_num,
        "q_num": q_num,
        "type": qtype,
        "question": title,
        "instruction": instruction,
        "options": options,
        "correct_answers": correct_indices
    })

print(f"Parsed {len(parsed_questions)} questions from CHỦ NGHĨA XH-KH.docx. Errors: {len(errors)}")

# Parse CNKHXH BAI6.docx (Bài 6)
import os
bai6_file = 'CNKHXH BAI6.docx'
if os.path.exists(bai6_file):
    doc_b6 = docx.Document(bai6_file)
    b6_raw = []
    cur_b6_q = None

    for p in doc_b6.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        m = re.match(r'^(Câu|CÂU)\s+(\d+)\s*[:.]?\s*(.*)', txt, re.IGNORECASE)
        if m:
            if cur_b6_q:
                b6_raw.append(cur_b6_q)
            cur_b6_q = {
                'q_num': int(m.group(2)),
                'title': txt,
                'prompt': m.group(3).strip(),
                'paragraphs': []
            }
        else:
            if cur_b6_q:
                cur_b6_q['paragraphs'].append(p)

    if cur_b6_q:
        b6_raw.append(cur_b6_q)

    for q in b6_raw:
        qid = len(parsed_questions) + 1
        q_num = q['q_num']
        title = q['title']
        paras = q['paragraphs']

        if q_num == 41:
            pairs = [
                {"left": "1. Biến đổi về quy mô gia đình", "right": "Gia đình hạt nhân"},
                {"left": "2. Biến đổi về chức năng tái sản xuất", "right": "Giảm số con"},
                {"left": "3. Biến đổi về quan hệ gia đình", "right": "Mâu thuẫn giữa các thế hệ"},
                {"left": "4. Biến đổi về chức năng kinh tế", "right": "Kinh tế hàng hóa"}
            ]
            options_pool = [
                "Gia đình hạt nhân",
                "Giảm số con",
                "Mâu thuẫn giữa các thế hệ",
                "Kinh tế hàng hóa"
            ]
            parsed_questions.append({
                "id": qid,
                "lesson": "Bài 6",
                "lesson_num": 6,
                "q_num": q_num,
                "type": "drag_drop_match",
                "question": title,
                "instruction": "Kéo hoặc chọn đáp án đúng ghép vào vế tương ứng",
                "pairs": pairs,
                "options_pool": options_pool
            })
            continue

        if q_num == 42:
            items = [
                {
                    "statement": "1. Hiện nay, độ bền vững của gia đình không bị chi phối bởi các mối quan hệ hòa hợp tình cảm giữa chồng và vợ, cha mẹ và con cái.",
                    "correct": False
                },
                {
                    "statement": "2. Hiện nay, kinh tế gia đình đang trở thành một bộ phận quan trọng trong nền kinh tế quốc dân.",
                    "correct": True
                },
                {
                    "statement": "3. Nội dung giáo dục trong gia đình hiện nay chỉ tập trung giáo dục đạo đức.",
                    "correct": False
                },
                {
                    "statement": "4. Giáo dục gia đình tiếp tục nhấn mạnh sự hy sinh của cá nhân cho cộng đồng.",
                    "correct": True
                }
            ]
            parsed_questions.append({
                "id": qid,
                "lesson": "Bài 6",
                "lesson_num": 6,
                "q_num": q_num,
                "type": "true_false",
                "question": title,
                "instruction": "Chọn Đúng hoặc Sai cho mỗi phát biểu sau",
                "items": items
            })
            continue

        # Single / Multiple Choice
        ans_text = ""
        opt_lines = []
        for p in paras:
            ptxt = p.text.strip()
            ans_m = re.match(r'^Đáp\s*án\s*(đúng)?\s*[:.]?\s*(.*)', ptxt, re.IGNORECASE)
            if ans_m:
                ans_text = ans_m.group(2).strip()
            else:
                for l in ptxt.split('\n'):
                    l = l.strip()
                    if l:
                        opt_lines.append(l)

        ans_letters = [x.strip().upper() for x in re.split(r'[,;\s]+', ans_text) if x.strip()]
        correct_indices = []
        options = []

        for idx_opt, line in enumerate(opt_lines):
            m = re.match(r'^([A-Z])\s*[\.\:\)]\s*(.*)', line)
            if m:
                letter = m.group(1).upper()
                content = m.group(2).strip()
                options.append(f"{letter}. {content}")
            else:
                letter = chr(65 + idx_opt)
                options.append(f"{letter}. {line}")

            if letter in ans_letters:
                correct_indices.append(idx_opt)

        is_multi_title = bool(re.search(r'chọn\s*[2345]\s*(đáp\s*án|phương\s*án)', title, re.IGNORECASE))
        is_multi = len(correct_indices) > 1 or is_multi_title
        qtype = "multiple_choice" if is_multi else "single_choice"

        instruction = "Chọn 1 đáp án đúng nhất"
        if qtype == "multiple_choice":
            num_ans = len(correct_indices) if len(correct_indices) > 1 else 2
            instruction = f"Chọn {num_ans} đáp án đúng"

        if len(correct_indices) == 0:
            errors.append(f"Bài 6 Q{q_num} ({qtype}): NO CORRECT ANSWER DETECTED! Options: {options}")

        parsed_questions.append({
            "id": qid,
            "lesson": "Bài 6",
            "lesson_num": 6,
            "q_num": q_num,
            "type": qtype,
            "question": title,
            "instruction": instruction,
            "options": options,
            "correct_answers": correct_indices
        })

print(f"Total parsed questions (All lessons including Bài 6): {len(parsed_questions)}. Errors: {len(errors)}")

# Save questions.json
with open('questions.json', 'w', encoding='utf-8') as f:
    json.dump(parsed_questions, f, ensure_ascii=False, indent=2)

# Also save questions-data.js
with open('questions-data.js', 'w', encoding='utf-8') as f:
    f.write(f"// Dữ liệu câu hỏi trắc nghiệm Chủ Nghĩa Xã Hội Khoa Học ({len(parsed_questions)} câu)\n")
    f.write("window.QUIZ_QUESTIONS = ")
    json.dump(parsed_questions, f, ensure_ascii=False, indent=2)
    f.write(";\n")

print("Saved questions.json and questions-data.js successfully!")
